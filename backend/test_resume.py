import io
import unittest
import pymupdf
from docx import Document
from starlette.testclient import TestClient
from app.main import app


class TestResumeEndpoints(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_demo_resume_endpoint(self):
        response = self.client.post("/api/resume/demo")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["candidate_name"], "Alex Johnson")
        self.assertIn("Python", data["skills"])
        self.assertIn("SQL", data["skills"])
        self.assertIn("Git", data["skills"])
        self.assertTrue(data["is_demo"])

    def test_pdf_upload_and_parsing(self):
        # Create an in-memory real PDF document
        pdf_doc = pymupdf.open()
        page = pdf_doc.new_page()
        resume_text = (
            "SARAH CONNOR\n"
            "sarah.connor@cyberdyne.io | +1 415-555-0199 | github.com/sconnor\n\n"
            "EDUCATION\n"
            "B.S. in Computer Engineering, Tech Institute 2023\n\n"
            "TECHNICAL SKILLS\n"
            "Python, FastAPI, Docker, PostgreSQL, React, AWS, Git, REST APIs\n\n"
            "EXPERIENCE\n"
            "Cloud Engineer at SkyNet (2023 - 2024)\n"
            "Containerized backend microservices and deployed Kubernetes clusters.\n\n"
            "PROJECTS\n"
            "Distributed Task Queue with Python and Redis\n"
        )
        page.insert_text((50, 50), resume_text, fontsize=11)
        pdf_bytes = pdf_doc.write()
        pdf_doc.close()

        files = {
            "file": ("sarah_resume.pdf", pdf_bytes, "application/pdf")
        }
        response = self.client.post("/api/resume/upload", files=files)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["candidate_name"], "SARAH CONNOR")
        self.assertEqual(data["email"], "sarah.connor@cyberdyne.io")
        self.assertIn("Python", data["skills"])
        self.assertIn("FastAPI", data["skills"])
        self.assertIn("Docker", data["skills"])
        self.assertIn("PostgreSQL", data["skills"])

    def test_docx_upload_and_parsing(self):
        # Create an in-memory real DOCX document
        doc = Document()
        doc.add_heading("DAVID MILLER", level=1)
        doc.add_paragraph("david.miller@example.com | 555-432-8765 | linkedin.com/in/dmiller")
        doc.add_heading("EDUCATION", level=2)
        doc.add_paragraph("Master of Science in Software Engineering, Metro University")
        doc.add_heading("TECHNICAL SKILLS", level=2)
        doc.add_paragraph("Java, Spring Boot, SQL, MySQL, Git, Linux, Microservices")
        
        docx_io = io.BytesIO()
        doc.save(docx_io)
        docx_bytes = docx_io.getvalue()

        files = {
            "file": ("david_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        response = self.client.post("/api/resume/upload", files=files)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["candidate_name"], "DAVID MILLER")
        self.assertEqual(data["email"], "david.miller@example.com")
        self.assertIn("Java", data["skills"])
        self.assertIn("Spring Boot", data["skills"])
        self.assertIn("MySQL", data["skills"])

    def test_unsupported_file_extension(self):
        files = {
            "file": ("malicious.exe", b"binarycontent", "application/octet-stream")
        }
        response = self.client.post("/api/resume/upload", files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("Unsupported file format", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
